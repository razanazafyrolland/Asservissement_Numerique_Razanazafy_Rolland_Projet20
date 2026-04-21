import numpy as np
import matplotlib.pyplot as plt
from scipy.integrate import odeint
import os

base_dir = os.path.dirname(os.path.abspath(__file__))
# Matrices du système
A = np.array([[0, 1],
              [-2, -3]])
B = np.array([[0],
              [1]])
C = np.array([[1, 0]])

# Gain de l'observateur
L = np.array([[5],
              [5]])

# Entrée (échelon)
def u(t):
    return 1

# Système réel
def system(x, t):
    dx = A @ x + B.flatten() * u(t)
    
    # Simulation de panne (à t > 5)
    if t > 5:
        dx[1] += 2   # perturbation = panne
    
    return dx

# Observateur
def observer(x_hat, t, x_real):
    y = C @ x_real
    y_hat = C @ x_hat
    dx_hat = A @ x_hat + B.flatten() * u(t) + L.flatten() * (y - y_hat)
    return dx_hat

# Simulation
t = np.linspace(0, 10, 1000)
x0 = [0, 0]
x_hat0 = [0, 0]

x = odeint(system, x0, t)

x_hat = np.zeros_like(x)

for i in range(len(t)-1):
    dt = t[i+1] - t[i]
    x_hat[i+1] = x_hat[i] + dt * observer(x_hat[i], t[i], x[i])

# Résidu
y = (C @ x.T).flatten()
y_hat = (C @ x_hat.T).flatten()
r = y - y_hat

# Seuil de détection
seuil = 0.05

# Détection (alarme)
alarme = np.abs(r) > seuil
# instant de première panne détectée
indice_panne = np.where(alarme == True)[0]

if len(indice_panne) > 0:
    t_panne = t[indice_panne[0]]
else:
    t_panne = None

print("Instant de panne détecté :", t_panne)

# --- Graphique ---
fig, ax = plt.subplots(3, 1, figsize=(9, 7))

# --- SORTIE ---
ax[0].plot(t, y, label="y réel")
ax[0].plot(t, y_hat, '--', label="y estimé")
ax[0].legend()
ax[0].set_title("Comparaison sortie")

# --- RÉSIDU ---
ax[1].plot(t, r, label="résidu")
ax[1].axhline(seuil, linestyle='--')
ax[1].axhline(-seuil, linestyle='--')

# 🔴 ZONE ROUGE DE PANNE
if t_panne is not None:
    ax[1].axvspan(t_panne, t[-1], color='red', alpha=0.2, label="zone panne")

    # 🔥 FLÈCHE D'ANNOTATION
    ax[1].annotate(
        f"             Panne détectée\n t = {t_panne:.2f} s",
        xy=(t_panne, r[np.argmin(np.abs(t - t_panne))]),
        xytext=(t_panne + 1, np.max(r)),
        arrowprops=dict(arrowstyle="->", color="red"),
        color="red"
    )

ax[1].legend()
ax[1].set_title("Résidu avec détection de panne")

# --- ALARME ---
ax[2].plot(t, alarme.astype(int), label="alarme (0/1)")
ax[2].legend()
ax[2].set_title("Signal de détection")

plt.tight_layout()
plt.show()
# 🔥 IMPORTANT : forcer rendu avant export
fig.canvas.draw()

# chemin image
os.path.join(base_dir, "Courbe_Enregistrer_Auto.png")

# sauvegarde propre
path_img = os.path.join(base_dir, "Courbe_Enregistrer_Auto.png")
fig.savefig(path_img, dpi=300, bbox_inches="tight")

plt.show()
plt.close(fig)
from docx import Document

doc = Document()

doc.add_heading("RAPPORT DE MINI-PROJET", 0)
doc.add_paragraph("\n")

doc.add_heading("Détection de panne par observateur de Luenberger", level=1)

doc.add_paragraph("\nAuteur : RAZANAZAFY Rolland")
doc.add_paragraph("Matière : Asservissement Numérique")
doc.add_paragraph("Filière : Mécatronique ")
doc.add_paragraph("Année : 2026")

doc.add_page_break()

doc.add_heading("Mini-projet : Détection de panne par observateur", 0)
doc.add_heading("1a. Schéma bloc du système", level=1)

doc.add_paragraph(
"""
        u(t)
         │
         ▼
   ┌──────────────┐
   │   SYSTÈME    │
   │   x' = Ax+Bu │
   └──────────────┘
         │ y(t)
         ▼
   ┌──────────────┐
   │ OBSERVATEUR  │
   │ de Luenberger|
   └──────────────┘
         │ ŷ(t)
         ▼
   Résidu r(t) = y - ŷ
"""
)
doc.add_heading("1b. Objectif", level=1)
doc.add_paragraph("Ce projet consiste à détecter une panne dans un système dynamique à l’aide d’un observateur de Luenberger.")

doc.add_heading("2. Modèle utilisé", level=1)
doc.add_paragraph("Système d’état : x' = Ax + Bu, y = Cx")

doc.add_heading("3. Principe", level=1)
doc.add_paragraph("On compare la sortie réelle et la sortie estimée. La différence (résidu) permet de détecter une panne.")

doc.add_heading("4. Résultat", level=1)
doc.add_paragraph("Avant t=5s : système normal. Après t=5s : apparition d’une panne détectée par augmentation du résidu.")
from docx.shared import Inches
if t_panne is not None:
    doc.add_paragraph(f"Panne détectée automatiquement à t = {t_panne:.2f} s")
else:
    doc.add_paragraph("Aucune panne détectée automatiquement")
table = doc.add_table(rows=4, cols=3)
table.style = "Table Grid"

# En-têtes
table.cell(0,0).text = "État"
table.cell(0,1).text = "Résidu r(t)"
table.cell(0,2).text = "Diagnostic"

# Ligne 1
table.cell(1,0).text = "Normal"
table.cell(1,1).text = "|r(t)| < seuil"
table.cell(1,2).text = "Système OK"

# Ligne 2
table.cell(2,0).text = "Début panne"
table.cell(2,1).text = "|r(t)| ≈ seuil"
table.cell(2,2).text = "Alerte"

# Ligne 3
table.cell(3,0).text = "Panne"
table.cell(3,1).text = "|r(t)| > seuil"
table.cell(3,2).text = "Défaut détecté"
from docx.shared import Inches
import os

path_img = os.path.join(base_dir, "Courbe_Enregistrer_Auto.png")
fig.savefig(path_img, dpi=300, bbox_inches="tight")
plt.show()

doc.add_heading("5. Résultats graphiques", level=1)
doc.add_picture(path_img, width=Inches(5))

doc.add_paragraph(
    f"Figure 1 : Évolution des sorties, du résidu et du signal de détection. "
    f"La panne est détectée à t = {t_panne:.2f} s par dépassement du seuil."
)
doc.add_heading("6. Conclusion", level=1)

doc.add_paragraph(
    "Ce projet a permis de mettre en œuvre un observateur de Luenberger pour la détection de panne. "
    "La comparaison entre la sortie réelle et estimée a permis de construire un résidu fiable. "
    "Les résultats montrent que le système est capable de détecter efficacement une anomalie. "
    "Cette méthode est largement utilisée dans les systèmes industriels de surveillance et de diagnostic."
)
path_word = os.path.join(base_dir, "Rapport_Enregistrer_Auto.docx")
doc.save(path_word)
print("✅ Rapport Word généré avec succès")